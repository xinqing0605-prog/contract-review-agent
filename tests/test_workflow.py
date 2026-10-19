"""检测规则与文件解析的单元测试（不依赖数据库和网络）。"""
import unittest

from backend.detector import analyze_text, sentence_span
from backend.ingest import extract_text

SEED = (
    "甲方在收到乙方发票并验收合格后九十日内支付全部款项，逾期付款不承担利息或其他责任。"
    "乙方在任何情况下不对间接损失、预期利润损失、数据损失承担责任。"
    "任一方未在合同期限届满前三十日书面通知终止的，本合同自动续期一年。"
    "双方对保密信息的保护义务在本合同终止后继续有效两年。"
)


class DetectorTests(unittest.TestCase):
    def test_detects_four_clause_types(self):
        clauses = analyze_text("doc-1", SEED)
        self.assertEqual(
            [clause["title"] for clause in clauses],
            ["付款与验收", "责任限制", "自动续期", "保密义务"],
        )

    def test_offsets_slice_back_to_original(self):
        for clause in analyze_text("doc-1", SEED):
            self.assertEqual(
                SEED[clause["start_offset"]: clause["end_offset"]], clause["original"]
            )

    def test_payment_with_long_terms_is_high_risk(self):
        clauses = analyze_text("doc-1", SEED)
        payment = [c for c in clauses if c["title"] == "付款与验收"][0]
        self.assertEqual(payment["risk"], "high")

    def test_payment_with_30_days_is_medium(self):
        clauses = analyze_text("doc-2", "甲方应在验收合格后三十日内支付服务费。")
        self.assertEqual(clauses[0]["risk"], "medium")

    def test_liability_cap_downgrades_risk(self):
        clauses = analyze_text("doc-3", "乙方赔偿责任总额不超过合同总金额的百分之三十。")
        self.assertEqual(clauses[0]["risk"], "medium")

    def test_force_majeure_is_not_a_liability_clause(self):
        clauses = analyze_text("doc-4", "如因不可抗力导致延期，双方互不承担责任。")
        self.assertEqual(clauses, [])

    def test_data_loss_is_not_data_processing(self):
        clauses = analyze_text("doc-5", "乙方不对数据损失承担责任。")
        self.assertEqual([c["title"] for c in clauses], ["责任限制"])

    def test_sentence_span_splits_on_chinese_period(self):
        text = "第一句。第二句命中关键词。第三句。"
        start = text.index("关键词")
        span = sentence_span(text, start, start + 3)
        self.assertEqual(text[span[0]: span[1]], "第二句命中关键词。")


class IngestTests(unittest.TestCase):
    def test_txt_utf8(self):
        source, text = extract_text(
            "合同.txt", "甲方应在验收合格后三十日内付款，逾期不承担利息。".encode("utf-8")
        )
        self.assertEqual(source, "txt")
        self.assertIn("三十日", text)

    def test_txt_gbk_fallback(self):
        source, text = extract_text(
            "合同.txt", "甲方应在验收合格后三十日内付款，逾期不承担利息。".encode("gb18030")
        )
        self.assertEqual(source, "txt")
        self.assertIn("三十日", text)

    def test_unsupported_suffix(self):
        with self.assertRaises(ValueError):
            extract_text("setup.exe", b"MZ")

    def test_too_short_after_parse(self):
        with self.assertRaises(ValueError):
            extract_text("空.txt", b"abc")

    def test_docx_paragraphs_and_tables(self):
        import io

        from docx import Document

        document = Document()
        document.add_paragraph("乙方不对间接损失承担责任。")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "期限"
        table.cell(0, 1).text = "自动续期一年"
        buffer = io.BytesIO()
        document.save(buffer)

        source, text = extract_text("合同.docx", buffer.getvalue())
        self.assertEqual(source, "docx")
        self.assertIn("间接损失", text)
        self.assertIn("自动续期一年", text)


if __name__ == "__main__":
    unittest.main()
