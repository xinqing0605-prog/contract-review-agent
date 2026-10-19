"""文件导入：把上传的文件变成纯文本。

真实合同几乎都是 PDF / Word，纯文本粘贴只是兜底。解析失败要给出可读的原因，
而不是抛一个 500。
"""
from __future__ import annotations

import io
import re
from pathlib import Path

SUPPORTED_SUFFIXES = {".txt", ".md", ".docx", ".pdf"}
MAX_BYTES = 8 * 1024 * 1024
MIN_CHARS = 20


def _decode_text(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return payload.decode("utf-8", errors="replace")


PAGE_MARK_RE = re.compile(r"^\[\s*第\s*\d+\s*页\s*\]$")
HEADING_RE = re.compile(r"^(第[一二三四五六七八九十百零〇\d]+[条章节]|附件[一二三四五六七八九十]|\[表格\])")
SENTENCE_END_RE = re.compile(r"[。！？；：!?;:）】」』”\"']$")


def _reflow(text: str) -> str:
    """把 PDF 里按版面断开的行重新合并成段落。

    PDF 抽出来的文本是逐行的：一句合同条款可能被拆成 3~4 行。
    直接展示会变成一行行碎句，这里按标点判断段落边界：
    - 行尾是句末标点 → 换段
    - 当前行是「第X条」「附件」「[表格]」「[第 N 页]」→ 独立成段
    - 其余情况把下一行接上来
    """
    lines = [line.strip() for line in text.split("\n")]
    out: list[str] = []
    for line in lines:
        if not line:
            if out and out[-1] != "":
                out.append("")
            continue
        if PAGE_MARK_RE.match(line):
            out.append(line)
            continue

        if not any(item.strip() for item in out):
            # 首行（通常是标题）独立成段，不和正文合并
            out.append(line)
            continue

        previous = out[-1] if out else ""
        filled = sum(1 for item in out if item.strip())
        mergeable = (
            previous
            and previous != ""
            and filled >= 2
            # 短且不以句末标点结尾的行视为标题/小标题，不把正文并进去
            and not (len(previous) <= 30 and not SENTENCE_END_RE.search(previous))
            and not PAGE_MARK_RE.match(previous)
            and not HEADING_RE.match(previous)
            and not SENTENCE_END_RE.search(previous)
            and not HEADING_RE.match(line)
        )
        if mergeable:
            out[-1] = previous + line
        else:
            out.append(line)
    return "\n".join(out)


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _from_docx(payload: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise ValueError("服务端缺少 python-docx，无法解析 Word") from exc

    document = Document(io.BytesIO(payload))
    blocks: list[str] = [paragraph.text for paragraph in document.paragraphs]
    # 合同里的关键信息经常在表格里（付款节点、交付清单），不能只读段落
    for table in document.tables:
        blocks.append("[表格]")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n".join(block for block in blocks if block is not None)


def _from_pdf(payload: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ValueError("服务端缺少 pypdf，无法解析 PDF") from exc

    reader = PdfReader(io.BytesIO(payload))
    pages = []
    for index, page in enumerate(reader.pages, 1):
        content = page.extract_text() or ""
        if content.strip():
            pages.append(f"[第 {index} 页]\n{content}")
    return "\n".join(pages)


def extract_text(filename: str, payload: bytes) -> tuple[str, str]:
    """返回 (来源类型, 正文)。校验失败抛 ValueError，由路由转成 422。"""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持 {suffix or '未知'} 格式，请上传 {'/'.join(sorted(SUPPORTED_SUFFIXES))}")
    if not payload:
        raise ValueError("文件是空的")
    if len(payload) > MAX_BYTES:
        raise ValueError(f"文件超过 {MAX_BYTES // 1024 // 1024} MB")

    if suffix in {".txt", ".md"}:
        text = _decode_text(payload)
    elif suffix == ".docx":
        text = _from_docx(payload)
    else:
        # PDF 是逐行抽取的，先重排成段落再展示，否则会变成一行行碎句
        text = _reflow(_from_pdf(payload))

    text = _normalize(text)
    if len(text) < MIN_CHARS:
        raise ValueError(
            f"解析后只有 {len(text)} 个字符，可能是扫描件或加密文件。"
            "扫描版 PDF 需要先做 OCR。"
        )
    return suffix.lstrip("."), text
